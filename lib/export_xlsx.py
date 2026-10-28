"""Excel with mockups of the selected designs: picture of the cap | type | name | code.

The pictures are embedded at full resolution (~1600 px per side) and only displayed small in the cell,
so in Excel they can be enlarged without losing quality.
"""
import datetime, io, json, os
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Font, PatternFill
from PIL import Image

import brand, shapes

OUT = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), 'designs')
CELL_W = 340          # displayed picture width in the cell, px


def cap_type(spec):
    from gemini_design import cap_kind
    kind = cap_kind(spec)
    if kind == 'pu_kids':
        return 'PU детская'
    if kind == 'pu_adult':
        return 'PU взрослая'
    if kind == 'long':
        return 'Для длинных волос'
    if kind in shapes.KIDS:
        return 'Детская фигурная (' + shapes.SHAPES[kind]['ru'].split(': ', 1)[-1] + ')'
    return 'Детская' if spec.get('segment') == 'Детская' else 'Взрослая'


def picture(slug, sides, cap_image, way=0):
    """Side A (and B next to it for two-sided designs) on white, full resolution, in colourway `way`."""
    ims = [cap_image(slug, s, way=way).convert('RGBA') for s in sides]
    h = max(i.height for i in ims)
    gap = 60
    out = Image.new('RGB', (sum(i.width for i in ims) + gap * (len(ims) - 1), h), 'white')
    x = 0
    for i in ims:
        out.paste(i, (x, h - i.height), i)
        x += i.width + gap
    return out


def _ral(h):
    import ral
    r = ral.nearest(h)
    return f"{r['code']} {r['name_en']}"


def rows(slugs, design_info, colorways):
    """One row per colourway: (slug, info, way index, colourway)."""
    for slug in slugs:
        info = design_info(slug)
        if not info:
            continue
        for k, cw in enumerate(colorways(slug)):
            yield slug, info, k, cw


def build(slugs, design_info, cap_image, colorways, pictures=None):
    """Excel: picture | type | name | code | cap colour | inks, one row per cap colour + print.
    pictures: optional dict (slug, way) -> PIL image, reused for the "save as" package."""
    wb = Workbook()
    ws = wb.active
    ws.title = 'Мокапы'
    head = (('A', 50, 'Шапочка'), ('B', 24, 'Тип'), ('C', 30, 'Название'), ('D', 26, 'Код'), ('E', 26, 'Цвет шапочки (RAL)'), ('F', 40, 'Краски принта (RAL)'))
    for col, w, t in head:
        ws.column_dimensions[col].width = w
        c = ws[f'{col}1']
        c.value, c.font, c.fill = t, Font(bold=True, color='FFFFFF'), PatternFill('solid', fgColor='1F2D3D')
        c.alignment = Alignment(vertical='center')
    ws.freeze_panes = 'A2'
    row, keep = 2, []
    for slug, info, k, cw in rows(slugs, design_info, colorways):
        spec = json.load(open(os.path.join(OUT, slug, '_gemini', 'spec.json'), encoding='utf-8'))
        sides = info['sides'] or ['A']
        pic = (pictures or {}).get((slug, k)) or picture(slug, sides, cap_image, k)
        if pictures is not None:
            pictures[(slug, k)] = pic
        buf = io.BytesIO()
        pic.save(buf, 'PNG', optimize=True)
        keep.append(buf)
        xi = XLImage(buf)
        w = CELL_W * (1 if len(sides) == 1 else 1.9)
        xi.width, xi.height = round(w), round(w * pic.height / pic.width)   # display size only, the file keeps full resolution
        ws.add_image(xi, f'A{row}')
        if len(sides) > 1:
            ws.column_dimensions['A'].width = max(ws.column_dimensions['A'].width, 92)
        ws.row_dimensions[row].height = xi.height * 0.75 + 12
        inks = '\n'.join(_ral(h) for h in cw['inks']) if cw['inks'] else 'полноцветный принт'
        vals = (('B', cap_type(spec)), ('C', info['name'] + ('' if k == 0 else f' (вариант {k})')), ('D', slug),
                ('E', _ral(cw['cap'])), ('F', inks))
        for col, val in vals:
            c = ws[f'{col}{row}']
            c.value, c.alignment = val, Alignment(vertical='center', wrap_text=True)
        ws[f'B{row}'].font = Font(bold=True)
        row += 1
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue(), f'{brand.NAME}_mockups_{datetime.date.today().isoformat()}.xlsx'


def safe_name(s):
    import re
    return re.sub(r'[<>:"/\\|?*\x00-\x1f]+', ' ', s).strip().rstrip('.')[:80]


SKIP = ('_gemini', '_guide', '_checks.json', '_hero.png')


def package(slugs, design_info, cap_image, colorways, out_dir):
    """Folder for "save as": <root>/BRAND_mockups_<date>.xlsx + one folder per model (mockup, spec, SVG/, PNG/,
    colourways/<n>_<RAL>.png). Returns the manifest [{path, file}] (paths relative to the chosen folder)."""
    import shutil
    os.makedirs(out_dir, exist_ok=True)
    pictures = {}
    data, name = build(slugs, design_info, cap_image, colorways, pictures)
    open(os.path.join(out_dir, name), 'wb').write(data)
    files = [dict(path=name, file=os.path.join(out_dir, name))]
    for slug in slugs:
        info = design_info(slug)
        if not info:
            continue
        folder = safe_name(f"{info['name']} [{slug}]")
        src = os.path.join(OUT, slug)
        for root, dirs, fs in os.walk(src):
            rel = os.path.relpath(root, src)
            if rel.split(os.sep)[0] in SKIP:
                continue
            for f in fs:
                if f.startswith('_') or f.endswith('.tmp.png'):
                    continue
                r = f if rel == '.' else f'{rel}/{f}'.replace(os.sep, '/')
                files.append(dict(path=f'{folder}/{r}', file=os.path.join(root, f)))
        for k, cw in enumerate(colorways(slug)):
            pic = pictures.get((slug, k))
            if pic is None:
                continue
            fn = f"{k + 1}_{_ral(cw['cap']).split(' ')[0]}_{_ral(cw['cap']).split(' ')[1]}.png"
            p = os.path.join(out_dir, f'{slug}_{fn}')
            pic.save(p)
            files.append(dict(path=f'{folder}/colourways/{fn}', file=p))
    return files
