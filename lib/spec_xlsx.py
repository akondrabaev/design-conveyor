"""Technical specification of a design as an Excel workbook, in English and Chinese, colours in RAL Classic.

designs/<slug>/<slug>_spec.xlsx
  sheet "Spec 规格":   design info, cap colours and print inks (RAL code, EN/ZH name, swatch, ΔE), print parameters,
                        checks, file list
  sheet "Visual 视图": mockup and layout sheets
Name and concept are translated once by Gemini and cached in _gemini/spec.json ("i18n").
"""
import datetime, glob, json, os
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from PIL import Image

import brand, ral
from build import OUT
from core import GAP, MINW, MARGIN

SEGMENT = {'Детская': ('Kids', '儿童'), 'Женская': ('Women, long hair', '女款（长发）'), 'Унисекс': ('Unisex', '中性')}
ROLE = {'main': ('Main', '主色'), 'accent': ('Accent', '辅色'), 'detail': ('Detail', '细节色')}
GRADE = {'exact': ('exact match', '完全一致'), 'close': ('close', '接近'), 'approx': ('approximate - check on sample', '近似，需打样确认')}

TRANSLATE = """Translate the design name and concept of a printed silicone swim cap into English and Simplified Chinese
for a factory technical specification. Keep them short and literal.
Name (Russian): {name}
Concept (Russian): {concept}
Return JSON: {{"en": {{"name": "...", "concept": "..."}}, "zh": {{"name": "...", "concept": "..."}}}}"""


def translations(slug, spec):
    """EN/ZH name and concept, cached in spec.json."""
    if spec.get('i18n', {}).get('zh'):
        return spec['i18n']
    try:
        import gemini_gen as gg
        from google.genai import types
        r = gg.client().models.generate_content(
            model=gg.TEXT_MODEL, contents=TRANSLATE.format(name=spec.get('name', slug), concept=spec.get('concept', '')),
            config=types.GenerateContentConfig(response_mime_type='application/json', temperature=0))
        spec['i18n'] = json.loads(r.text)
        json.dump(spec, open(f'{OUT}/{slug}/_gemini/spec.json', 'w', encoding='utf-8'), ensure_ascii=False, indent=1)
    except Exception:   # offline / no key: the spec is still written, with the slug as the name
        return {'en': {'name': slug, 'concept': ''}, 'zh': {'name': slug, 'concept': ''}}
    return spec['i18n']


def _lum(h):
    h = h.lstrip('#')
    r, g, b = (int(h[i:i + 2], 16) / 255 for i in (0, 2, 4))
    return 0.2126 * r + 0.7152 * g + 0.0722 * b


# ---------- styles ----------
THIN = Side(style='thin', color='C8C8C8')
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEAD = PatternFill('solid', fgColor='1F2D3D')
SUB = PatternFill('solid', fgColor='EEF1F4')
WRAP = Alignment(wrap_text=True, vertical='top')


def _fill(hex_):
    return PatternFill('solid', fgColor=hex_.lstrip('#').upper())


class Sheet:
    def __init__(self, ws):
        self.ws, self.r = ws, 1

    def title(self, text, size=16):
        c = self.ws.cell(self.r, 1, text)
        c.font = Font(bold=True, size=size)
        self.r += 2 if size >= 14 else 1

    def section(self, text):
        self.r += 1
        c = self.ws.cell(self.r, 1, text)
        c.font = Font(bold=True, size=12, color='1F2D3D')
        self.r += 1

    def header(self, cols):
        for i, t in enumerate(cols, 1):
            c = self.ws.cell(self.r, i, t)
            c.font, c.fill, c.alignment, c.border = Font(bold=True, color='FFFFFF'), HEAD, WRAP, BOX
        self.r += 1

    def row(self, vals, fills=None, bold_first=False):
        for i, v in enumerate(vals, 1):
            c = self.ws.cell(self.r, i, brand.text(v))   # brand token -> BRAND_NAME
            c.alignment, c.border = WRAP, BOX
            if bold_first and i == 1:
                c.font, c.fill = Font(bold=True), SUB
            if fills and i in fills:
                c.fill = _fill(fills[i])
        self.r += 1


def write(slug):
    od = f'{OUT}/{slug}'
    spec = json.load(open(f'{od}/_gemini/spec.json', encoding='utf-8'))
    tr = translations(slug, spec)
    if spec.get('print') == 'full_colour':
        return write_pu(slug, spec, tr)
    checks = json.load(open(f'{od}/_checks.json', encoding='utf-8')) if os.path.exists(f'{od}/_checks.json') else {}
    qrep = json.load(open(f'{od}/_gemini/quantize_report.json', encoding='utf-8')) if os.path.exists(f'{od}/_gemini/quantize_report.json') else {}
    sides = [s for s in ('A', 'B') if os.path.exists(f'{od}/PNG/{slug}_side-{s}_print_transparent.png')]
    import gemini_design, shapes
    kind = gemini_design.cap_kind(spec)
    long = kind == 'long'
    sizes = shapes.sizes_for(kind, spec.get('segment'))
    main = next(iter(sizes.values()))
    cap_mm = main['w'] * 10
    mm = cap_mm / 1000
    seg = SEGMENT.get(spec.get('segment'), ('Unisex', '中性'))
    inks = spec['inks']
    logo_ink = max(range(len(inks)), key=lambda i: abs(_lum(inks[i]['hex']) - _lum(spec['base']['hex'])))
    snapped = spec.get('palette') == 'RAL Classic'

    wb = Workbook()
    ws = wb.active
    ws.title = 'Spec 规格'
    for col, w in zip('ABCDEFGHIJ', (30, 14, 22, 16, 12, 10, 24, 16, 44, 44)):
        ws.column_dimensions[col].width = w
    S = Sheet(ws)
    S.title(f'{brand.NAME} swim cap - technical specification / {brand.NAME} 泳帽 - 技术规格')

    sides_en = ('2 sides: A (wearer\'s left, with logo) + B (wearer\'s right, ' +
                ('mirror image of A, no logo)' if spec.get('side_b') == 'mirror' else 'own print, no logo)')) if len(sides) > 1 else \
        '1 side: A (wearer\'s left, with logo); side B is plain'
    sides_zh = ('双面：A面（佩戴者左侧，带标志）+ B面（佩戴者右侧，' + ('A面的镜像图案，无标志）' if spec.get('side_b') == 'mirror' else '独立图案，无标志）')) \
        if len(sides) > 1 else '单面：A面（佩戴者左侧，带标志），B面无印刷'
    import gemini_design as _gd
    lmode = _gd.logo_mode(spec)
    if kind in shapes.FRONT:   # front/back halves (caps with ears / horns)
        sides_en = ('2 sides: A = front (the face) + B = back of the head' if len(sides) > 1 else
                    '1 side: A = front (the face); the back of the cap is plain')
        sides_zh = '双面：A面=正面（脸部）+ B面=背面' if len(sides) > 1 else '单面：A面=正面（脸部），背面无印刷'
    info = [
        ('Design / 设计', f"{tr['en']['name']}\n{tr['zh']['name']}"),
        ('Design code / 设计编号', slug),
        ('Concept / 设计理念', f"{tr['en']['concept']}\n{tr['zh']['concept']}"),
        ('Product / 产品', ('Silicone swim cap for long hair (extra volume at the back)\n长发硅胶泳帽（后部加大容量）' if long else
                          f"Silicone {shapes.SHAPES[kind]['en'].lower()}" + (' (the print covers the moulded parts too)' if kind in shapes.KIDS else '') +
                          f"\n{shapes.SHAPES[kind]['zh']}（硅胶" + ('，立体造型部分同样印刷）' if kind in shapes.KIDS else '）'))
                         + f"\n{seg[0]} / {seg[1]}"),
        ('Printing / 印刷', f"Screen printing, {len(inks)} inks, no overprint between inks\n丝网印刷，{len(inks)}色，各颜色之间不叠印"),
        ('Print sides / 印刷面', f'{sides_en}\n{sides_zh}'),
        ('Logo / 标志', ('No logo / 无标志' if lmode == 'none' else
                        'BRAND lettering is part of the artwork (already in the stencils) / BRAND字样为图案的一部分（已包含在网版中）' if lmode == 'styled' else
                        (f"Small BRAND wordmark, {'bottom centre under the face' if kind in shapes.FRONT else 'lower back corner of side A'}, ink {logo_ink + 1}\n"
                         f"小号BRAND字标，{'位于脸部下方正中' if kind in shapes.FRONT else '位于A面后下角'}，使用第{logo_ink + 1}号印刷色") if lmode == 'small' else
                        (f"BRAND logo on side A, front lower corner (near the left temple), printed with ink {logo_ink + 1}\n"
                         f"BRAND标志位于A面前下角（靠近左太阳穴），使用第{logo_ink + 1}号印刷色"))),
        ('Colour standard / 颜色标准', 'RAL Classic. Screen colours are approximate: approve the cap and ink colours on physical samples.\n'
                                    '以RAL经典色卡为准。屏幕颜色仅供参考，泳帽和油墨颜色须以实物样品确认。'),
        ('Date / 日期', datetime.date.today().isoformat()),
    ]
    for k, v in info:
        S.row([k, v], bold_first=True)
        ws.merge_cells(start_row=S.r - 1, start_column=2, end_row=S.r - 1, end_column=10)
        ws.row_dimensions[S.r - 1].height = 15 * (str(v).count('\n') + 1) + 4

    note = ('' if snapped else ' (ΔE2000 = difference between the design colour and the RAL colour / 设计颜色与RAL颜色的色差)')
    S.section('Cap colours / 泳帽颜色' + note)
    S.header(['Variant / 款式', 'RAL', 'RAL name (EN)', 'RAL 名称 (中文)', 'RAL colour / 色样', 'ΔE2000', 'Match / 匹配', 'Design HEX / 设计色值', '', ''])
    caps = [('Main cap colour / 主色', spec['base'])] + [(f'Alternative {i + 1} / 备选{i + 1}', v) for i, v in enumerate(spec.get('variants', [])[:2])]
    for label, c in caps:
        r = ral.nearest(c['hex'])
        g = GRADE[ral.grade(r['de'])]
        S.row([label, r['code'], r['name_en'], r['name_zh'], '', r['de'], f'{g[0]} / {g[1]}', c['hex'], '', ''], fills={5: r['hex']})

    import colorways, copy
    cw = colorways.ensure(copy.deepcopy(spec))
    S.section('Colourways: same stencils, ink colours per cap colour / 配色方案：同一网版，每种帽色对应的油墨颜色')
    S.header(['Colourway / 配色', 'Cap RAL / 帽色', 'Sample / 色样'] + sum([[f'Ink {i + 1} RAL / 第{i + 1}色', 'Sample / 色样'] for i in range(len(spec['inks']))], [])
             + [''] * max(0, 10 - 3 - 2 * len(spec['inks'])))
    ways = [('Main / 主配色', spec['base']['hex'], [i['hex'] for i in spec['inks']])] + \
           [(f'Alternative {n + 1} / 备选{n + 1}', v['hex'], [i['hex'] for i in v['inks']]) for n, v in enumerate(cw.get('variants', [])[:2])]
    for label, cap_hex, ink_hex in ways:
        cells, fills = [label, f"{ral.nearest(cap_hex)['code']} {ral.nearest(cap_hex)['name_en']} / {ral.nearest(cap_hex)['name_zh']}", ''], {3: cap_hex}
        for h in ink_hex:
            r = ral.nearest(h)
            cells += [f"{r['code']} {r['name_en']} / {r['name_zh']}", '']
            fills[len(cells)] = r['hex']
        S.row(cells + [''] * max(0, 10 - len(cells)), fills=fills)
        ws.row_dimensions[S.r - 1].height = 34

    S.section('Print inks (main colourway) / 印刷颜色（主配色）')
    S.header(['Ink / 颜色', 'RAL', 'RAL name (EN)', 'RAL 名称 (中文)', 'RAL colour / 色样', 'ΔE2000', 'Match / 匹配',
              'Coverage, side A / A面覆盖率', 'Stencil SVG / 网版矢量文件', 'Stencil PNG / 网版图片'])
    for i, k in enumerate(inks):
        r = ral.nearest(k['hex'])
        role = ROLE.get(k.get('role'), ('', ''))
        cov = qrep.get('A', {}).get('inks', {}).get(str(i + 1), {}).get('coverage_pct')
        fname = k.get('file', f'ink{i + 1}').replace('.svg', '')
        svgs = '\n'.join(os.path.basename(p) for p in sorted(glob.glob(f'{od}/SVG/{slug}_side-*_stencil{i + 1}_*.svg')))
        pngs = '\n'.join(os.path.basename(p) for p in sorted(glob.glob(f'{od}/PNG/{slug}_side-*_stencil{i + 1}_*.png')))
        g = GRADE[ral.grade(r['de'])]
        S.row([f"{i + 1}. {role[0]} / {role[1]}" + (' + logo / 标志' if i == logo_ink else ''), r['code'], r['name_en'], r['name_zh'], '',
               r['de'], f'{g[0]} / {g[1]}', f'{cov} %' if cov is not None else '',
               svgs or f'SVG/{slug}_side-A_stencil{i + 1}_{fname}.svg', pngs], fills={5: r['hex']})
        ws.row_dimensions[S.r - 1].height = 15 * max(1, len(sides)) + 4

    S.section('Sizes / 尺寸 (flat cap; width without moulded parts / 平放尺寸，宽度不含立体造型部分)')
    S.header(['Size / 尺寸', 'Weight / 重量', 'Width / 宽度', 'Height / 高度', 'Rim width / 帽口宽度',
              'Artwork scale / 图稿比例', 'Print scale vs main size / 相对主尺寸缩放', '', '', ''])
    for i, (key, z) in enumerate(sizes.items()):
        S.row([('Main / 主尺寸: ' if i == 0 else '') + z['en'] + f" / {z['g']}克", f"{z['g']} g",
               f"{z['w'] * 10:.0f} mm", f"{z['h'] * 10:.0f} mm", f"{z['rim'] * 10:.0f} mm" if z.get('rim') else '-',
               f"1000 px = {z['w'] * 10:.0f} mm", f"{z['w'] / main['w'] * 100:.1f} %", '', '', ''])

    S.section('Print parameters (main size) / 印刷参数（主尺寸）')
    params = [
        ('Artwork scale / 图稿比例', f'1000 px = {cap_mm:.0f} mm in SVG; PNG files are 3x (3000 px) / SVG中1000像素={cap_mm:.0f}毫米；PNG为3倍（3000像素）'),
        ('Print margin from edge / 边缘留白', f'{MARGIN * mm:.0f} mm ({MARGIN} px)'),
        ('Minimum line width / 最小线宽', f'{MINW * mm:.1f} mm ({MINW} px)'),
        ('Gap between inks / 颜色间隙', f'{GAP * mm:.1f} mm ({GAP} px), inks never overlap / 各颜色不重叠'),
    ]
    if len(sides) > 1:
        params.append(('Side B / B面', ('Printed on the other side of the flat cap, artwork already oriented for it (front on the right).\n'
                                        '印于平放泳帽的另一面，图稿方向已调整（前部在右侧）。')))
    for k, v in params:
        S.row([k, v], bold_first=True)
        ws.merge_cells(start_row=S.r - 1, start_column=2, end_row=S.r - 1, end_column=10)
        ws.row_dimensions[S.r - 1].height = 15 * (v.count('\n') + 1) + 4

    S.section('Automatic checks / 自动检查')
    S.header(['Side / 面', 'Inks used / 颜色数', 'Inside print area / 在印刷区内', 'Min gap, px / 最小间隙', 'Result / 结果', '', '', '', ''])
    for s in sides:
        c = checks.get(s, {})
        ok = c.get('inks_ok') and c.get('inside_safe_ok') and c.get('gap_ok', True)
        S.row([s, c.get('inks_used'), 'yes / 是' if c.get('inside_safe_ok') else 'NO / 否', c.get('min_gap_between_inks_px'),
               'OK / 合格' if ok else 'CHECK / 需检查', '', '', '', ''])

    S.section('Files / 文件')
    S.header(['File / 文件', 'Content / 内容', '', '', '', '', '', '', '', ''])
    desc = {
        '_mockup.png': 'Visual of the cap in its colour variants / 泳帽各颜色效果图',
        '_spec.xlsx': 'This specification / 本技术规格',
        '_print_vector.svg': 'Whole print, all inks, vector / 完整印刷图案（全部颜色，矢量）',
        '_print_transparent.png': 'Whole print on transparent background / 透明背景完整印刷图案',
        '_layout.png': 'Print placed on the flat cap template with the print area / 图案在平放泳帽模板上的位置及印刷区域',
        '_stencil': 'One stencil (screen) per ink, black = printed area / 每色一张网版，黑色为印刷区域',
    }
    files = [f for f in sorted(os.listdir(od)) if not f.startswith('_') and os.path.isfile(f'{od}/{f}')] + \
            [f'{sub}/{f}' for sub in ('SVG', 'PNG') if os.path.isdir(f'{od}/{sub}') for f in sorted(os.listdir(f'{od}/{sub}'))]
    for f in files:
        S.row([f, next((v for k, v in desc.items() if k in f), ''), '', '', '', '', '', '', '', ''])
        ws.merge_cells(start_row=S.r - 1, start_column=2, end_row=S.r - 1, end_column=10)
    ws.column_dimensions['A'].width = 58

    return _visuals_and_save(wb, od, slug, sides)

def _visuals_and_save(wb, od, slug, sides):
    """Sheet with the mockup and the layouts, then save."""
    vs = wb.create_sheet('Visual 视图')
    vs['A1'] = 'Mockup / 效果图'
    vs['A1'].font = Font(bold=True, size=13)
    row = 2
    for path, title in [(f'{od}/{slug}_mockup.png', None)] + \
            [(f'{od}/PNG/{slug}_side-{s}_layout.png', f'Layout, side {s} / {s}面排版') for s in sides]:
        if not os.path.exists(path):
            continue
        if title:
            vs.cell(row, 1, title).font = Font(bold=True, size=13)
            row += 1
        im = Image.open(path)
        w = 900
        h = round(im.height * w / im.width)
        thumb = path + '.xlsx.png'
        im.convert('RGB').resize((w, h)).save(thumb)
        xi = XLImage(thumb)
        vs.add_image(xi, f'A{row}')
        row += h // 20 + 3
    out = f'{od}/{slug}_spec.xlsx'
    wb.save(out)
    for t in set(glob.glob(f'{od}/**/*.xlsx.png', recursive=True)):   # temporary previews embedded above
        os.remove(t)
    return out



def pu_logo_text(lg, spec, size):
    """EN / ZH description of the logo layer of a PU print."""
    if not lg:
        return 'No logo / 无标志' if spec.get('upload') else 'BRAND lettering is part of the artwork / BRAND字样为图案的一部分'
    mm = lg['s'] * size['w'] * 10
    if lg.get('kind') == 'styled':
        return (f"BRAND lettering in the style of the print, about {mm:.0f} mm wide, part of the print file\n"
                f"与图案同风格的BRAND字样，宽约{mm:.0f}毫米，已包含在印刷文件中")
    code = ral.nearest(lg['hex'])['code']
    return (f"BRAND logo ({'wordmark' if lg['kind'] == 'adult' else 'with waves'}), {code}, about {mm:.0f} mm wide, part of the print file\n"
            f"BRAND标志（{'字标' if lg['kind'] == 'adult' else '带浪花'}），{code}，宽约{mm:.0f}毫米，已包含在印刷文件中")


def write_pu(slug, spec, tr):
    """PU cap: one full-colour print (heat transfer), no inks and stencils."""
    import shapes
    od = f'{OUT}/{slug}'
    kind = spec['shape']
    size = shapes.main_size(kind)
    checks = json.load(open(f'{od}/_checks.json', encoding='utf-8')).get('A', {})
    pmm = checks.get('print_mm', ['?', '?'])
    wb = Workbook()
    ws = wb.active
    ws.title = 'Spec 规格'
    for col, w in zip('ABCDEFGHIJ', (34, 14, 22, 16, 12, 10, 24, 16, 20, 20)):
        ws.column_dimensions[col].width = w
    S = Sheet(ws)
    S.title(f'{brand.NAME} PU swim cap - technical specification / {brand.NAME} PU泳帽 - 技术规格')
    sh = shapes.SHAPES[kind]
    info = [
        ('Design / 设计', f"{tr['en']['name']}\n{tr['zh']['name']}"),
        ('Design code / 设计编号', slug),
        ('Concept / 设计理念', f"{tr['en']['concept']}\n{tr['zh']['concept']}"),
        ('Product / 产品', f"{sh['en']}: PU-coated fabric, stitched seams\n{sh['zh']}：PU涂层面料，缝线结构"),
        ('Cap size (flat) / 泳帽尺寸（平放）', f"{size['w'] * 10:.0f} x {size['h'] * 10:.0f} mm / {size['w'] * 10:.0f}×{size['h'] * 10:.0f}毫米"),
        ('Printing / 印刷', 'Full-colour print (heat transfer / DTF, or as agreed), one side: A = wearer\'s left side panel\n'
                           '全彩印刷（热转印/DTF或协商工艺），单面：A面=佩戴者左侧面板'),
        ('Print size / 印刷尺寸', f"{pmm[0]} x {pmm[1]} mm, file PNG/{slug}_print.png, 300 dpi, transparent background\n"
                                f"{pmm[0]}×{pmm[1]}毫米，文件PNG/{slug}_print.png，300dpi，透明背景"),
        ('Print position / 印刷位置', 'All-over: the whole side of the cap, cut to the outline, see the layout sheet\n整面印刷：泳帽整个侧面，按轮廓裁切，见排版图'
                                      if spec.get('pu_layout') == 'allover' else
                                      'Centred on the side panel, see the layout sheet\n位于侧面板中央，见排版图'),
        ('Logo / 标志', pu_logo_text(spec.get('pu_logo'), spec, size)),
        ('Colour standard / 颜色标准', 'Cap colour: RAL Classic. Print: full colour, match the PNG file; approve on a physical sample.\n'
                                    '泳帽颜色以RAL经典色卡为准；印刷为全彩，以PNG文件为准，须以实物样品确认。'),
        ('Date / 日期', datetime.date.today().isoformat()),
    ]
    for k, v in info:
        S.row([k, v], bold_first=True)
        ws.merge_cells(start_row=S.r - 1, start_column=2, end_row=S.r - 1, end_column=10)
        ws.row_dimensions[S.r - 1].height = 15 * (str(v).count('\n') + 1) + 4

    S.section('Cap colours / 泳帽颜色')
    S.header(['Variant / 款式', 'RAL', 'RAL name (EN)', 'RAL 名称 (中文)', 'RAL colour / 色样', 'ΔE2000', 'Match / 匹配', 'Design HEX / 设计色值', '', ''])
    caps = [('Main cap colour / 主色', spec['base'])] + [(f'Alternative {i + 1} / 备选{i + 1}', v) for i, v in enumerate(spec.get('variants', [])[:2])]
    for label, c in caps:
        r = ral.nearest(c['hex'])
        g = GRADE[ral.grade(r['de'])]
        S.row([label, r['code'], r['name_en'], r['name_zh'], '', r['de'], f'{g[0]} / {g[1]}', c['hex'], '', ''], fills={5: r['hex']})

    S.section('Main print colours (reference only) / 主要印刷颜色（仅供参考）')
    S.header(['Colour / 颜色', 'Share / 占比', 'HEX', 'Nearest RAL / 最接近RAL', 'Sample / 色样', 'ΔE2000', '', '', '', ''])
    for i, k in enumerate(spec.get('inks', [])):
        r = ral.nearest(k['hex'])
        S.row([f'{i + 1}', f"{k.get('share', '')} %", k['hex'], f"{r['code']} {r['name_en']} / {r['name_zh']}", '', r['de'], '', '', '', ''],
              fills={5: k['hex']})

    S.section('Files / 文件')
    S.header(['File / 文件', 'Content / 内容', '', '', '', '', '', '', '', ''])
    desc = {'_mockup.png': 'Visual of the cap in its colour variants / 泳帽各颜色效果图', '_spec.xlsx': 'This specification / 本技术规格',
            '_print.png': 'PRINT FILE: artwork at print size, 300 dpi, transparent / 印刷文件：实际尺寸、300dpi、透明背景',
            '_print_transparent.png': 'Print placed on the flat cap template (preview) / 图案在平放模板上的位置（预览）',
            '_layout.png': 'Layout with the print area and dimensions / 带印刷区域和尺寸的排版图'}
    files = [f for f in sorted(os.listdir(od)) if not f.startswith('_') and os.path.isfile(f'{od}/{f}')] + \
            [f'PNG/{f}' for f in sorted(os.listdir(f'{od}/PNG'))]
    for f in files:
        S.row([f, next((v for k, v in desc.items() if f.endswith(k)), ''), '', '', '', '', '', '', '', ''])
        ws.merge_cells(start_row=S.r - 1, start_column=2, end_row=S.r - 1, end_column=10)
    ws.column_dimensions['A'].width = 52
    return _visuals_and_save(wb, od, slug, ['A'])
